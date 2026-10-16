"""Bank statement parsers -> RawTxn. Pure functions (no DB) so they are trivial to test.

Supported: delimited CSV / XLSX with automatic header mapping (HDFC/ICICI/SBI/Axis/Kotak style),
SWIFT MT940, ISO-20022 camt.053, and payment-gateway settlement CSVs.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from ..money import parse_amount, to_minor


@dataclass
class RawTxn:
    txn_date: date
    amount_minor: int                 # always positive
    direction: str                    # credit | debit
    narration: str = ""
    value_date: date | None = None
    bank_ref: str | None = None
    balance_minor: int | None = None
    currency: str = "INR"
    fee_minor: int = 0
    utr_hint: str | None = None
    extra: dict = field(default_factory=dict)


class StatementError(ValueError):
    pass


_DATE_FORMATS = ["%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y", "%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%Y-%m-%d",
                 "%d.%m.%Y", "%d/%b/%Y", "%m/%d/%Y"]


def parse_date(raw) -> date | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    s = re.sub(r"\s+\d{1,2}:\d{2}(:\d{2})?.*$", "", s)  # drop time part
    for f in _DATE_FORMATS:
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


# header synonyms -> canonical column
_SYN = {
    "date": ["txn date", "transaction date", "date", "tran date", "post date", "posting date", "trans date", "txn dt"],
    "value_date": ["value date", "value dt", "val date", "value dt."],
    "narration": ["narration", "description", "particulars", "transaction remarks", "remarks", "transaction particulars",
                  "details", "transaction details"],
    "ref": ["chq./ref.no.", "chq/ref no", "ref no", "reference no", "cheque no", "chq no", "ref no./cheque no.",
            "chq./ref no.", "reference", "utr", "utr no", "utr number", "instrument no"],
    "debit": ["withdrawal amt.", "withdrawal", "debit", "dr", "withdrawals", "debit amount", "withdrawal amount", "dr amount"],
    "credit": ["deposit amt.", "deposit", "credit", "cr", "deposits", "credit amount", "deposit amount", "cr amount"],
    "amount": ["amount", "txn amount", "transaction amount"],
    "drcr": ["dr/cr", "cr/dr", "type", "txn type", "transaction type"],
    "balance": ["closing balance", "balance", "running balance", "available balance", "bal"],
}


def _norm_h(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _map_headers(row: list) -> dict[str, int] | None:
    idx: dict[str, int] = {}
    for i, h in enumerate(row):
        nh = _norm_h(h)
        for canon, syns in _SYN.items():
            if nh in syns and canon not in idx:
                idx[canon] = i
    if "date" in idx and "narration" in idx and (("debit" in idx and "credit" in idx) or "amount" in idx):
        return idx
    return None


def _rows_to_txns(rows: list[list]) -> list[RawTxn]:
    hdr_i, idx = None, None
    for i, r in enumerate(rows[:60]):
        m = _map_headers(r)
        if m:
            hdr_i, idx = i, m
            break
    if idx is None:
        raise StatementError("Could not locate a header row with date, narration and amount/debit/credit columns")
    out: list[RawTxn] = []
    for n, r in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):
        def cell(k):
            i = idx.get(k)
            return r[i] if i is not None and i < len(r) else None

        d = parse_date(cell("date"))
        if d is None:
            continue  # footers / blank / opening-balance rows
        debit = parse_amount(_s(cell("debit"))) if "debit" in idx else None
        credit = parse_amount(_s(cell("credit"))) if "credit" in idx else None
        direction, amt = None, None
        if credit and credit > 0:
            direction, amt = "credit", credit
        elif debit and debit > 0:
            direction, amt = "debit", debit
        elif "amount" in idx:
            a = parse_amount(_s(cell("amount")))
            if a:
                t = _s(cell("drcr")).upper()
                raw = _s(cell("amount")).upper()
                is_debit = t.startswith(("D", "W")) or "DR" in raw or a < 0
                direction, amt = ("debit" if is_debit else "credit"), abs(a)
        if not amt:
            continue
        out.append(RawTxn(
            txn_date=d, value_date=parse_date(cell("value_date")), amount_minor=amt, direction=direction,
            narration=re.sub(r"\s+", " ", _s(cell("narration"))).strip(), bank_ref=_s(cell("ref")) or None,
            balance_minor=parse_amount(_s(cell("balance"))) if "balance" in idx else None,
            extra={"row": n},
        ))
    return out


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def parse_csv(data: bytes) -> list[RawTxn]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    return _rows_to_txns(rows)


def parse_xlsx(data: bytes) -> list[RawTxn]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.active
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    return _rows_to_txns(rows)


# ------------------------------------------------------------------ MT940
def parse_mt940(data: bytes) -> list[RawTxn]:
    text = data.decode("utf-8", errors="replace")
    ccy = "INR"
    if m := re.search(r":60[FM]:[CD]\d{6}([A-Z]{3})", text):
        ccy = m.group(1)
    out: list[RawTxn] = []
    blocks = re.split(r"(?m)^(?=:61:)", text)
    for b in blocks:
        if not b.startswith(":61:"):
            continue
        m = re.match(r":61:(\d{6})(\d{4})?(R?[CD])[A-Z]?([\d,]+)([A-Z][A-Z0-9]{3})([^\r\n]*)", b)
        if not m:
            continue
        vd, md, cd, amt, _typ, refline = m.groups()
        vdate = datetime.strptime(vd, "%y%m%d").date()
        tdate = vdate
        if md:
            try:
                tdate = vdate.replace(month=int(md[:2]), day=int(md[2:]))
            except ValueError:
                pass
        info = ""
        if m86 := re.search(r":86:(.*?)(?=\n:6[01]|\n:62|\n-\}|\Z)", b, re.S):
            info = re.sub(r"\s*\n\s*", " ", m86.group(1)).strip()
            info = re.sub(r"\?\d{2}", " ", info)  # German-style structured subfields
        out.append(RawTxn(
            txn_date=tdate, value_date=vdate, amount_minor=to_minor(amt.replace(",", ".")),
            direction="debit" if cd.endswith("D") else "credit", narration=re.sub(r"\s+", " ", info),
            bank_ref=refline.split("//")[0].strip() or None, currency=ccy,
        ))
    if not out:
        raise StatementError("No :61: statement lines found in MT940 file")
    return out


# ------------------------------------------------------------------ camt.053
def parse_camt053(data: bytes) -> list[RawTxn]:
    from lxml import etree

    root = etree.fromstring(data)
    ns = {"c": root.nsmap.get(None) or "urn:iso:std:iso:20022:tech:xsd:camt.053.001.02"}
    out: list[RawTxn] = []
    for ntry in root.iterfind(".//c:Ntry", ns):
        amt_el = ntry.find("c:Amt", ns)
        if amt_el is None:
            continue
        cd = (ntry.findtext("c:CdtDbtInd", namespaces=ns) or "").upper()
        bk = ntry.findtext("c:BookgDt/c:Dt", namespaces=ns) or ntry.findtext("c:BookgDt/c:DtTm", namespaces=ns)
        vl = ntry.findtext("c:ValDt/c:Dt", namespaces=ns)
        parts = [ntry.findtext("c:AddtlNtryInf", namespaces=ns) or ""]
        parts += [e.text or "" for e in ntry.iterfind(".//c:RmtInf/c:Ustrd", ns)]
        parts += [e.text or "" for e in ntry.iterfind(".//c:RltdPties/c:Dbtr/c:Nm", ns)]
        utr = ntry.findtext(".//c:Refs/c:EndToEndId", namespaces=ns) or ntry.findtext(".//c:Refs/c:UETR", namespaces=ns)
        out.append(RawTxn(
            txn_date=parse_date(bk[:10]) or date.today(), value_date=parse_date(vl[:10]) if vl else None,
            amount_minor=to_minor(amt_el.text), direction="credit" if cd == "CRDT" else "debit",
            narration=" ".join(p for p in parts if p).strip(), currency=amt_el.get("Ccy", "INR"),
            utr_hint=utr if utr and utr != "NOTPROVIDED" else None,
            bank_ref=ntry.findtext("c:AcctSvcrRef", namespaces=ns),
        ))
    if not out:
        raise StatementError("No <Ntry> entries found in camt.053 file")
    return out


# ------------------------------------------------------------------ gateway settlements
def parse_gateway_csv(data: bytes) -> list[RawTxn]:
    """Settlement report with one row per payment: gross, fee, tax, net, settlement utr, order/notes."""
    text = data.decode("utf-8-sig", errors="replace")
    rd = csv.DictReader(io.StringIO(text))
    if not rd.fieldnames:
        raise StatementError("Empty gateway file")
    f = {_norm_h(h): h for h in rd.fieldnames}

    def pick(*names):
        for n in names:
            if n in f:
                return f[n]
        return None

    c_gross = pick("amount", "gross amount", "payment amount", "gross")
    c_fee = pick("fee", "fees", "mdr", "gateway fee")
    c_tax = pick("tax", "gst", "tax on fee", "gst on fee")
    c_date = pick("settled at", "settlement date", "date", "settled on")
    c_utr = pick("settlement utr", "utr", "settlement id")
    c_ref = pick("order id", "order_id", "notes", "reference", "invoice", "description", "invoice no", "payment id")
    c_id = pick("payment id", "id", "transaction id")
    if not (c_gross and c_date):
        raise StatementError("Gateway file needs at least amount and settlement date columns")
    out: list[RawTxn] = []
    for r in rd:
        gross = parse_amount(r.get(c_gross))
        d = parse_date(r.get(c_date))
        if not gross or not d:
            continue
        fee = (parse_amount(r.get(c_fee)) or 0) if c_fee else 0
        tax = (parse_amount(r.get(c_tax)) or 0) if c_tax else 0
        ref = (r.get(c_ref) or "").strip() if c_ref else ""
        out.append(RawTxn(
            txn_date=d, amount_minor=gross - fee - tax, fee_minor=fee + tax, direction="credit",
            narration=f"GATEWAY SETTLEMENT {ref}".strip(), bank_ref=(r.get(c_id) or "").strip() or None,
            utr_hint=(r.get(c_utr) or "").strip() or None,
            extra={"gross_minor": gross, "gateway_ref": ref, "fee_only": fee, "tax_on_fee": tax},
        ))
    return out


def detect_and_parse(filename: str, data: bytes) -> tuple[str, list[RawTxn]]:
    name = (filename or "").lower()
    head = data[:4000].decode("utf-8", errors="ignore")
    if name.endswith((".xlsx", ".xlsm")):
        return "xlsx", parse_xlsx(data)
    if name.endswith(".xml") or head.lstrip().startswith("<?xml") or "<Document" in head:
        return "camt053", parse_camt053(data)
    if ":61:" in head or ":20:" in head and ":25:" in head:
        return "mt940", parse_mt940(data)
    return "csv", parse_csv(data)
