"""Remittance advice extraction (email bodies, PDF text, portal exports).

`RuleExtractor` is deterministic and always available. `LLMExtractor` is optional, redacts PII before
anything leaves the process, and ALWAYS falls back to rules on any failure. Matching never *depends* on
an LLM: extracted values are only evidence that the scorer weighs, and every extracted invoice reference
must resolve against real open items before it counts.
"""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from typing import Protocol

from ..money import parse_amount

_AMT = r"(?:₹|Rs\.?|INR)?\s*\d{1,3}(?:,\d{2,3})*(?:\.\d{1,2})?|\d+\.\d{1,2}"
_TOTAL_AMT = r"(?:₹|Rs\.?|INR)?\s*(?:\d{1,3}(?:,\d{2,3})+(?:\.\d{1,2})?|\d+\.\d{1,2}|\d{4,})"   # a total must look like money (grouping, decimals or 4+ digits)
_DATE_LIKE = re.compile(r"^\d{1,2}[-/.](?:\d{1,2}|[A-Za-z]{3})[-/.]\d{2,4}$")
_UTR_LABELED = re.compile(r"(?i)\b(?:UTR|RRN|ref(?:erence)?\s*(?:no|number|#)?)\s*[:#\-]?\s*([A-Z0-9]{10,30})\b")
_PAN = re.compile(r"\b[A-Z]{5}\d{4}[A-Z]\b")
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE = re.compile(r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{9}(?!\d)")
_LONGNUM = re.compile(r"(?<![\w.,])\d{11,20}(?![\w.,])")


def redact(text: str) -> str:
    """Best-effort PII masking before any external call. Invoice numbers and amounts are preserved."""
    t = _PAN.sub("[PAN]", text)
    t = _EMAIL.sub("[EMAIL]", t)
    t = _PHONE.sub("[PHONE]", t)
    return _LONGNUM.sub("[ACCT]", t)


class RemittanceExtractor(Protocol):
    name: str

    def extract(self, text: str) -> dict: ...


def _empty() -> dict:
    return {"payer_name": None, "utr": None, "total_minor": None, "tds_total_minor": None, "lines": [], "confidence": 0.0}


class RuleExtractor:
    name = "rules"

    def extract(self, text: str) -> dict:
        tbl = parse_table(text)
        out = self._extract_text(text)
        if tbl and tbl["lines"] and len(tbl["lines"]) >= len(out["lines"]):   # a recognised header row beats guessing at columns
            out["lines"] = tbl["lines"]
            out["total_minor"] = tbl["total_minor"] if tbl["total_minor"] is not None else out["total_minor"]
            out["tds_total_minor"] = tbl["tds_total_minor"] if tbl["tds_total_minor"] is not None else out["tds_total_minor"]
            out["confidence"] = 0.8
            out["meta"] = {"parse": "table"}
        else:
            out["meta"] = {"parse": "text"}
        return finalise(out)

    def _extract_text(self, text: str) -> dict:
        out = _empty()
        if not text:
            return out
        if m := _UTR_LABELED.search(text):
            out["utr"] = m.group(1).upper()
        if m := re.search(r"(?im)^\s*(?:from|payer|customer|vendor|company)\s*[:\-]\s*(.+)$", text):
            out["payer_name"] = m.group(1).strip()[:200]
        if m := re.search(rf"(?i)(?:total\s+(?:payment|paid|amount)|net\s+(?:amount\s+)?paid|amount\s+paid|payment\s+amount|total\s+net)[^\d\n]{{0,12}}({_TOTAL_AMT})", text):
            out["total_minor"] = parse_amount(m.group(1))
        if m := re.search(rf"(?i)total\s+tds[^\d\n]{{0,12}}({_TOTAL_AMT})", text):
            out["tds_total_minor"] = parse_amount(m.group(1))

        lines = []
        for raw in text.splitlines():
            ln = raw.strip()
            if not ln or re.match(r"(?i)^(total|grand|sub\s*total|net\s+paid|amount\s+paid)", ln):
                continue
            if re.search(r"(?i)\b(utr|rrn)\b", ln) or (out["utr"] and out["utr"] in ln.upper()):
                continue
            ref = self._first_doc_ref(ln)
            if not ref:
                continue
            tds = None
            if m := re.search(rf"(?i)\bTDS\b\D{{0,10}}({_AMT})", ln):
                tds = parse_amount(m.group(1))
            ded = None
            if m := re.search(rf"(?i)\b(?:deduction|discount|short(?:age)?|penalty|rate\s+diff\w*)\b\D{{0,10}}({_AMT})", ln):
                ded = parse_amount(m.group(1))
            net = None
            if m := re.search(rf"(?i)\b(?:net|paid|payment|pay)\b\D{{0,10}}({_AMT})", ln):
                net = parse_amount(m.group(1))
            stripped = re.sub(rf"(?i)\b(?:TDS|deduction|discount|net|paid|payment|pay)\b\D{{0,10}}(?:{_AMT})", "", ln)
            stripped = stripped.replace(ref, " ")
            stripped = re.sub(r"\b\d{1,2}[-/.](?:\d{1,2}|[A-Za-z]{3})[-/.]\d{2,4}\b", " ", stripped)  # drop dates
            amounts = []
            for m in re.finditer(rf"(\(\s*)?(?<![\w/-])({_AMT})(?![\w/])", stripped):
                if not re.search(r"[.,]", m.group(2)):
                    continue  # bare integers are day numbers / counts, not money
                v = parse_amount(m.group(2))
                if v:
                    amounts.append(-v if m.group(1) else v)
            gross = amounts[0] if amounts else None
            if net is None and gross is not None:
                net = gross - (tds or 0) - (ded or 0) if (tds or ded) else gross
            lines.append({"ref": ref, "gross_minor": gross, "amount_minor": net, "tds_minor": tds, "deduction_minor": ded,
                          "kind": "credit_note" if re.search(r"(?i)\b(CN|credit\s*note)\b", ln) else "invoice"})
        out["lines"] = lines
        out["confidence"] = 0.6 if lines else (0.3 if out["utr"] else 0.0)
        return out

    @staticmethod
    def _first_doc_ref(line: str) -> str | None:
        for i, tok in enumerate(re.split(r"[\s,;|]+", line)):
            tok = tok.strip("()[]:")
            if len(tok) < 4 or not any(c.isdigit() for c in tok) or _DATE_LIKE.match(tok):
                continue
            if re.fullmatch(r"[\d,]+(\.\d{1,2})?", tok) and not (i == 0 and "," not in tok and "." not in tok):
                continue  # money
            if re.fullmatch(r"\d{10,}", tok) and re.search(r"(?i)utr|rrn", line):
                continue
            if any(c.isalpha() for c in tok) or i == 0:
                return tok
        return None


class LLMExtractor:
    """Optional. Calls the Anthropic Messages API over HTTPS with redacted text. UNTESTED against the live API
    in this repo's CI (tests inject a fake transport); validate with your own key before relying on it."""

    name = "llm"

    def __init__(self, api_key: str, model: str, fallback: RemittanceExtractor | None = None, transport=None):
        self.api_key, self.model = api_key, model
        self.fallback = fallback or RuleExtractor()
        self.transport = transport

    PROMPT = (
        "Extract payment remittance details from the text. Reply with ONLY JSON: "
        '{"payer_name":str|null,"utr":str|null,"total":number|null,"tds_total":number|null,'
        '"lines":[{"ref":str,"gross":number|null,"paid":number|null,"tds":number|null,"deduction":number|null}]}. '
        "Amounts are in rupees. Never invent values; use null when absent."
    )

    def extract(self, text: str) -> dict:
        try:
            body = {"model": self.model, "max_tokens": 1500,
                    "messages": [{"role": "user", "content": f"{self.PROMPT}\n\n---\n{redact(text)[:12000]}"}]}
            if self.transport:
                raw = self.transport(body)
            else:
                import httpx

                from ..config import check_egress

                check_egress("https://api.anthropic.com")   # raises in data-residency strict mode unless explicitly allow-listed

                r = httpx.post("https://api.anthropic.com/v1/messages", json=body, timeout=30,
                               headers={"x-api-key": self.api_key, "anthropic-version": "2023-06-01"})
                r.raise_for_status()
                raw = r.json()["content"][0]["text"]
            data = json.loads(raw[raw.index("{"): raw.rindex("}") + 1])
            return self._validate(data, text)
        except Exception:
            return self.fallback.extract(text)

    @staticmethod
    def _m(v) -> int | None:
        return None if v is None else int(round(float(v) * 100))

    def _validate(self, d: dict, original: str) -> dict:
        out = _empty()
        out["payer_name"] = str(d["payer_name"])[:200] if d.get("payer_name") else None
        out["utr"] = str(d["utr"]).upper()[:30] if d.get("utr") else None
        out["total_minor"], out["tds_total_minor"] = self._m(d.get("total")), self._m(d.get("tds_total"))
        for ln in d.get("lines") or []:
            ref = str(ln.get("ref") or "").strip()
            if not ref or ref.upper() not in original.upper():  # reject hallucinated references
                continue
            out["lines"].append({"ref": ref, "gross_minor": self._m(ln.get("gross")), "amount_minor": self._m(ln.get("paid")),
                                 "tds_minor": self._m(ln.get("tds")), "deduction_minor": self._m(ln.get("deduction")), "kind": "invoice"})
        out["confidence"] = 0.75 if out["lines"] else 0.3
        out["meta"] = {"parse": "llm"}
        return finalise(out)


# ------------------------------------------------------------------------------------------------ tables, reconciliation, cascade
_HDR_ORDER = [   # leftmost-first alternation: more specific phrases come first ("Net Amount" before "Amount", "Invoice Amount" before "Invoice")
    ("net", r"net\s*(?:amount|paid|payable|payment)?|(?:amount\s*)?paid(?:\s*amount)?|payment(?:\s*amount)?|payable"),
    ("tds", r"tds(?:\s*(?:amount|deducted))?|tax\s*deducted(?:\s*at\s*source)?"),
    ("ded", r"deductions?|discount|rebate|short(?:age)?|penalty|rate\s*diff\w*|debit\s*note"),
    ("gross", r"(?:invoice|bill|document|gross)\s*(?:amount|value|amt)|gross|amount|amt"),
    ("date", r"(?:invoice|doc(?:ument)?|bill)?\s*date"),
    ("ref", r"(?:invoice|inv\.?|bill|document|doc\.?)\s*(?:no\.?|number|#)|reference(?:\s*(?:no\.?|number))?|ref\.?(?:\s*no\.?)?|invoice|bill|document"),
]
_HDR_RE = re.compile(r"\b(?:" + "|".join(f"(?P<{k}>{v})" for k, v in _HDR_ORDER) + r")\b", re.I)
_SPLIT = re.compile(r"\t")
_MONEY = re.compile(r"(?<![\w/.\-,])(\()?\s*-?\s*(?:₹|Rs\.?|INR)?\s*(\d{1,3}(?:,\d{2,3})+(?:\.\d{1,2})?|\d+\.\d{1,2}|\d{4,})(\))?(?![\w/\-]|\.\d)")
_AMOUNT_KINDS = ("gross", "tds", "ded", "net")


def _header_kinds(line: str) -> list[str]:
    return [m.lastgroup for m in _HDR_RE.finditer(line) if m.lastgroup]


def _cell_kind(cell: str) -> str | None:
    c = cell.strip()
    m = _HDR_RE.fullmatch(c) if c else None
    return m.lastgroup if m else None


def _money_tokens(text: str) -> list[int]:
    out = []
    for m in _MONEY.finditer(text):
        if re.search(r"[-/.]\d", text[max(0, m.start() - 1):m.start() + 1]):
            continue
        v = parse_amount(m.group(2))
        if v is None:
            continue
        out.append(-v if (m.group(1) and m.group(3)) or text[m.start():m.start(2)].strip().startswith("-") else v)
    return out


def _cell_money(cell: str) -> int | None:
    """A cell under an amount header holds a number whatever its size (a bare 500 is ₹500 here, unlike in running text)."""
    c = cell.strip().replace("₹", "").replace("Rs.", "").replace("INR", "").strip()
    m = re.fullmatch(r"(\()?\s*(-)?\s*(\d[\d,]*(?:\.\d{1,2})?)\s*(\))?", c)
    if not m:
        return None
    v = parse_amount(m.group(3))
    return None if v is None else (-v if (m.group(1) and m.group(4)) or m.group(2) else v)


def _first_ref(line: str) -> str | None:
    return RuleExtractor._first_doc_ref(line)


def _assign(amount_kinds: list[str], vals: list[int]) -> dict | None:
    """Map money tokens to the header's amount columns. Exact count: in order. One short: a blank cell (usually TDS or deductions) —
    accept only if exactly one assignment satisfies gross - tds - deductions = net, otherwise refuse to guess."""
    if len(vals) == len(amount_kinds):
        return dict(zip(amount_kinds, vals))
    if len(vals) == 1 and ("net" in amount_kinds or "gross" in amount_kinds):
        return {("net" if "net" in amount_kinds else "gross"): vals[0]}
    if len(vals) == len(amount_kinds) - 1 and "net" in amount_kinds and "gross" in amount_kinds:
        ok = []
        for drop in (k for k in amount_kinds if k in ("tds", "ded")):
            kinds = [k for k in amount_kinds if k != drop]
            m = dict(zip(kinds, vals))
            if m["gross"] - m.get("tds", 0) - m.get("ded", 0) == m["net"]:
                ok.append(m)
        if len(ok) == 1:
            return ok[0]
    return None


def parse_table(text: str) -> dict | None:
    """Find a header row naming an invoice reference and an amount column, then read the rows beneath it. Works on tab-separated sheets
    (cells are positional, blanks preserved) and on OCR/PDF text where columns are only separated by whitespace (arithmetic must agree).
    Returns None when there is no recognisable table."""
    lines = text.splitlines()
    for hi, raw in enumerate(lines):
        kinds = _header_kinds(raw)
        amount_kinds = [k for k in kinds if k in _AMOUNT_KINDS]
        if "ref" not in kinds or not ({"gross", "net"} & set(amount_kinds)):
            continue
        tabbed = "\t" in raw
        cells = [c.strip() for c in raw.split("\t")] if tabbed else []
        ckinds = [_cell_kind(c) for c in cells]
        rows, total_net, total_tds, skipped = [], None, None, 0
        for r in lines[hi + 1:]:
            if not r.strip():
                if rows:
                    break
                continue
            if re.match(r"(?i)^\s*(grand\s+)?total\b", r.strip()):
                vals = _money_tokens(r)
                m = _assign(amount_kinds, vals) if vals else None
                if m:
                    total_net, total_tds = m.get("net", m.get("gross")), m.get("tds")
                elif vals:
                    total_net = vals[-1]
                break
            m = None
            if tabbed:
                parts = [c.strip() for c in r.split("\t")]
                if len(parts) == len(cells):
                    ref = parts[ckinds.index("ref")] if "ref" in ckinds else ""
                    vals = {k: _cell_money(parts[i]) for i, k in enumerate(ckinds) if k in _AMOUNT_KINDS}
                    m = {k: v for k, v in vals.items() if v is not None}
                    if not ref or not any(ch.isdigit() for ch in ref):
                        continue
            if m is None:
                ref = _first_ref(r)
                if not ref:
                    continue
                m = _assign(amount_kinds, _money_tokens(r.replace(ref, " ", 1)))
            if not m:
                skipped += 1
                continue
            gross, tds, ded, net = m.get("gross"), m.get("tds"), m.get("ded"), m.get("net")
            if net is None and gross is not None:
                net = gross - (tds or 0) - (ded or 0)
            if gross is None and net is not None:
                gross = net + (tds or 0) + (ded or 0)
            if net is None:
                continue
            rows.append({"ref": ref, "gross_minor": gross, "amount_minor": net, "tds_minor": tds, "deduction_minor": ded,
                         "kind": "credit_note" if re.search(r"(?i)\b(CN|credit\s*note)\b", ref) or net < 0 else "invoice"})
        if rows:
            out = {"lines": rows, "total_minor": total_net, "tds_total_minor": total_tds}
            if skipped:
                out["skipped_rows"] = skipped
            return out
    return None


def finalise(out: dict) -> dict:
    """Self-checks that decide how far the result can be trusted: do the lines add up to the stated total?"""
    meta = out.setdefault("meta", {})
    lines = out.get("lines") or []
    warnings = list(meta.get("warnings") or [])
    total = out.get("total_minor")
    nets = [ln["amount_minor"] for ln in lines if ln.get("amount_minor") is not None]
    reconciled = None
    if total is not None and lines and len(nets) == len(lines):
        diff = total - sum(nets)
        reconciled = abs(diff) <= 100
        if not reconciled:
            warnings.append(f"Lines add up to {sum(nets) / 100:,.2f} but the advice states {total / 100:,.2f} (difference {diff / 100:,.2f}); a line may be missing or misread")
    for ln in lines:
        g, n, t, d = ln.get("gross_minor"), ln.get("amount_minor"), ln.get("tds_minor") or 0, ln.get("deduction_minor") or 0
        if g is not None and n is not None and (t or d) and abs(g - t - d - n) > 100:
            warnings.append(f"{ln['ref']}: gross - TDS - deductions does not equal the net amount")
    conf = out.get("confidence", 0.0)
    if reconciled is True:
        conf += 0.15
    elif reconciled is False:
        conf -= 0.25
    if any("does not equal the net" in w for w in warnings):
        conf -= 0.1
    meta["reconciled"], meta["warnings"] = reconciled, warnings
    out["confidence"] = round(max(0.0, min(conf, 0.95)), 2)
    return out


class CascadeExtractor:
    """Local rules first. The LLM runs only when the rules are not confident (or their lines do not add up), only when one is configured
    and the residency policy allows the call, and its answer is used only if it validates AND is more trustworthy than the rules' answer.
    Matching never depends on this: the output is evidence."""

    def __init__(self, rules: RemittanceExtractor, llm: RemittanceExtractor | None = None, threshold: float = 0.75):
        self.rules, self.llm, self.threshold = rules, llm, threshold
        self.name = "rules"

    def extract(self, text: str) -> dict:
        r = self.rules.extract(text)
        self.name = "rules"
        if not self.llm or r["confidence"] >= self.threshold:
            return r
        try:
            l = self.llm.extract(text)
        except Exception:
            return r
        if l.get("meta", {}).get("parse") != "llm":
            return r   # the LLM call failed and fell back to rules: nothing new to weigh
        self.name = "rules+llm"
        if l["confidence"] > r["confidence"] and l["lines"]:
            self.name = "llm"
            l["meta"]["rules_confidence"] = r["confidence"]
            return l
        r["meta"]["llm_consulted"] = True
        return r


def get_extractor() -> RemittanceExtractor:
    from ..config import egress_allowed, settings

    if settings.anthropic_api_key and egress_allowed("api.anthropic.com"):
        return CascadeExtractor(RuleExtractor(), LLMExtractor(settings.anthropic_api_key, settings.llm_model), settings.llm_fallback_below)
    return CascadeExtractor(RuleExtractor())


# ------------------------------------------------------------------------------------------------ uploads
@dataclass
class UploadText:
    text: str
    method: str = "text"                # text | ocr
    ocr_confidence: float | None = None
    pages: int | None = None
    warnings: list[str] = field(default_factory=list)


def _rows_to_text(rows) -> str:
    return "\n".join("\t".join("" if c is None else str(c).strip() for c in r) for r in rows)


def extract_upload(filename: str, data: bytes) -> UploadText:
    from . import ocr

    n = (filename or "").lower()
    if n.endswith(".pdf"):
        from pypdf import PdfReader

        pages = PdfReader(io.BytesIO(data)).pages
        txt = "\n".join((p.extract_text() or "") for p in pages)
        if len(re.sub(r"\s", "", txt)) >= 40 * max(1, min(len(pages), 1)):
            return UploadText(txt, "text", pages=len(pages))
        if not ocr.available():       # an image-only PDF we cannot read must not look like an empty remittance
            raise ocr.OcrUnavailable("This PDF is a scan (no text layer) and OCR is not installed on this server")
        t, c, k = ocr.pdf_to_text(data)
        return UploadText(t, "ocr", c, k, _ocr_warnings(c, k, len(pages)))
    if n.endswith(ocr.IMAGE_EXT):
        t, c = ocr.image_to_text(data)
        return UploadText(t, "ocr", c, 1, _ocr_warnings(c, 1, 1))
    if n.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).active
        return UploadText(_rows_to_text(ws.iter_rows(values_only=True)))
    if n.endswith(".csv"):
        import csv

        raw = data.decode("utf-8-sig", errors="replace")
        return UploadText(_rows_to_text(csv.reader(io.StringIO(raw))))
    if n.endswith(".eml"):
        import email
        from email import policy

        msg = email.message_from_bytes(data, policy=policy.default)
        body = msg.get_body(preferencelist=("plain", "html"))
        head = f"From: {msg.get('From', '')}\nSubject: {msg.get('Subject', '')}\n"
        return UploadText(head + (body.get_content() if body else ""))
    return UploadText(data.decode("utf-8", errors="replace"))


def _ocr_warnings(conf: float, read: int, total: int) -> list[str]:
    w = []
    if conf < 60:
        w.append(f"Low OCR quality ({conf:.0f}%): check every figure against the original image")
    if read < total:
        w.append(f"Only the first {read} of {total} pages were scanned")
    return w


def text_from_upload(filename: str, data: bytes) -> str:
    return extract_upload(filename, data).text
